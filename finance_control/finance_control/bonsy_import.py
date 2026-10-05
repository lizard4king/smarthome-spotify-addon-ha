"""Local, idempotent Bonsy XLSX receipt import; never creates ledger postings."""
from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from .classification import (
    _audit,
    _cash_withdrawal_evidence,
    _document_link_rejected,
    auto_link_documents,
    document_recipient_matches,
    link_document,
    normalize_counterparty,
)
from .core import money
from .document_intake import cache_document_source
from .import_preview import MAX_BYTES, _headers, inspect_workbook, outside_repository
from .transfer_corrections import effective_transfer_id, source_declares_transfer

MAX_ROWS = 20_000
_EXCLUDED_BONSY_WARNING = 'source_excluded_bonsy'


def _booking_minutes(value):
    """Read valid ISO-T / day-first-T purchase timestamps; ignore other formats."""
    result = set()
    if not isinstance(value, str):
        return result
    time_pattern = r'T(?P<hour>[0-9]{2}):(?P<minute>[0-9]{2})(?::(?P<second>[0-9]{2}))?(?![0-9:])'
    for date_pattern in (
            r'(?<!\d)(?P<year>[0-9]{4})-(?P<month>[0-9]{2})-(?P<day>[0-9]{2})',
            r'(?<!\d)(?P<day>[0-9]{2})-(?P<month>[0-9]{2})-(?P<year>[0-9]{4})'):
        for match in re.finditer(date_pattern + time_pattern, value):
            parts = {key: int(number or 0) for key, number in match.groupdict().items()}
            try:
                # Validate components only; retain the literal purchase clock time.
                purchase_time = datetime(**parts, tzinfo=UTC)
            except ValueError:
                continue
            result.add(purchase_time.strftime('%Y-%m-%dT%H:%M'))
    return result


def _bonsy_recipient_matches(vendor, counterparty, description, title):
    """Keep title-independent Bonsy matching limited to exact booking recipient evidence."""
    if document_recipient_matches(vendor, counterparty, description, title):
        return True
    if not isinstance(vendor, str) or not isinstance(counterparty, str) or not isinstance(description, str):
        return False
    vendor_key = normalize_counterparty(vendor)
    if len(vendor_key) < 3:
        return False
    booking_key = normalize_counterparty(f'{counterparty} {description}')
    return re.search(rf'(?<!\w){re.escape(vendor_key)}(?!\w)', booking_key) is not None


def auto_link_timestamped_receipts(store, ids=None):
    """Link a Bonsy receipt when amount and embedded purchase minute are unique."""
    params = ['%"source_excluded_bonsy"%', '%"duplicate_source_document"%', '%"not_invoice_like"%']
    query = '''SELECT r.entry_id,r.document_id,r.occurred_at,r.total,d.vendor,d.title
               FROM bonsy_receipts r
               JOIN classification_documents d ON d.id=r.document_id
               WHERE d.kind='invoice' AND d.status='confirmed'
                 AND d.warnings NOT LIKE ? AND d.warnings NOT LIKE ? AND d.warnings NOT LIKE ?
                 AND CAST(r.total AS REAL)>0
                 AND NOT EXISTS (SELECT 1 FROM classification_document_links l
                                 WHERE l.document_id=r.document_id)
                 AND NOT EXISTS (SELECT 1 FROM bonsy_cash_allocations a
                                 WHERE a.entry_id=r.entry_id)'''
    if ids is not None:
        ids = sorted(set(ids))
        if not ids:
            return {'linked': [], 'rejected': [], 'checked': 0}
        query += ' AND r.document_id IN (' + ','.join('?' for _ in ids) + ')'
        params.extend(ids)
    rows = store.db.execute(query + ' ORDER BY r.document_id', params).fetchall()
    linked, rejected = [], []
    for receipt in rows:
        occurred = datetime.fromisoformat(receipt['occurred_at'])
        minute = occurred.strftime('%Y-%m-%dT%H:%M')
        if store.db.execute(
                '''SELECT count(*) FROM bonsy_receipts r
                   JOIN classification_documents d ON d.id=r.document_id
                   WHERE substr(r.occurred_at,1,16)=? AND r.total=?
                     AND d.warnings NOT LIKE ? AND d.warnings NOT LIKE ? AND d.warnings NOT LIKE ?''',
                (minute, receipt['total'], '%"source_excluded_bonsy"%',
                 '%"duplicate_source_document"%', '%"not_invoice_like"%')).fetchone()[0] != 1:
            rejected.append({'id': receipt['document_id'],
                             'reasons': ['unique_receipt_timestamp_amount_required']})
            continue
        end = (occurred.date() + timedelta(days=7)).isoformat()
        minute_candidates = []
        for candidate in store.db.execute(
                '''SELECT t.account_id,t.external_id,t.amount,t.category,t.transfer_id,
                          COALESCE(c.counterparty,'') AS counterparty,
                          COALESCE(c.description,'') AS description
                   FROM transactions t
                   LEFT JOIN transaction_context c USING(account_id,external_id)
                   LEFT JOIN transfer_correction_members m
                     ON m.account_id=t.account_id AND m.external_id=t.external_id
                   WHERE t.date BETWEEN ? AND ? AND t.currency='EUR'
                     AND CAST(t.amount AS REAL)<0 AND (t.transfer_id IS NULL OR t.transfer_id='')
                     AND m.pair_id IS NULL
                     AND NOT EXISTS (SELECT 1 FROM classification_document_links l
                                     JOIN classification_documents d ON d.id=l.document_id
                                     WHERE l.account_id=t.account_id AND l.external_id=t.external_id
                                       AND d.kind='invoice')''',
                (occurred.date().isoformat(), end)):
            if (effective_transfer_id(store, candidate) or source_declares_transfer(candidate)
                    or -money(candidate['amount']) != money(receipt['total'])):
                continue
            if _cash_withdrawal_evidence(
                    store, candidate['account_id'], candidate['external_id']):
                continue
            if not _bonsy_recipient_matches(
                    receipt['vendor'], candidate['counterparty'], candidate['description'],
                    receipt['title']):
                continue
            if _document_link_rejected(
                    store, candidate['account_id'], candidate['external_id'], receipt['document_id']):
                continue
            booking_minutes = _booking_minutes(candidate['description'])
            if minute in booking_minutes:
                minute_candidates.append(candidate)
        if len(minute_candidates) != 1:
            rejected.append({'id': receipt['document_id'],
                             'reasons': ['unique_timestamped_transaction_required']})
            continue
        candidate = minute_candidates[0]
        link = link_document(store, {
            'account_id': candidate['account_id'],
            'external_id': candidate['external_id'],
            'document_id': receipt['document_id'], 'confirmed': True,
        })['link']
        linked.append(link)
    return {'linked': linked, 'rejected': rejected, 'checked': len(rows)}


def _text(value, name, maximum=500, *, optional=False):
    if value is None or (isinstance(value, str) and not value.strip()):
        if optional:
            return None
        raise ValueError(f'missing_bonsy_{name}')
    if isinstance(value, str) and value.startswith('='):
        raise ValueError('bonsy_formulas_not_supported')
    result = str(value).strip()
    if len(result) > maximum or any(ord(char) < 32 and char not in '\t\n\r' for char in result):
        raise ValueError(f'invalid_bonsy_{name}')
    return result


def _day(value):
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, datetime.min.time())
    elif isinstance(value, str) and not value.startswith('='):
        result = datetime.fromisoformat(value.strip())
    else:
        raise ValueError('invalid_bonsy_date')
    if result.date() < date(2026, 1, 1):
        raise ValueError('bonsy_before_2026')
    return result.isoformat(timespec='seconds')


def _decimal(value, name, *, optional=False):
    if value is None or str(value).strip() == '':
        if optional:
            return None
        raise ValueError(f'missing_bonsy_{name}')
    if isinstance(value, (bool, float)):
        value = str(value)
    try:
        result = Decimal(str(value).replace(',', '.'))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f'invalid_bonsy_{name}') from error
    if not result.is_finite():
        raise ValueError(f'invalid_bonsy_{name}')
    return format(result, 'f')


def _blank_source_value(value):
    return value is None or isinstance(value, str) and not value.strip()


def _valid_preview_date(value):
    if isinstance(value, (datetime, date)):
        return True
    if not isinstance(value, str) or value.startswith('='):
        return False
    try:
        datetime.fromisoformat(value)
    except ValueError:
        return False
    return True


def _valid_preview_amount(value):
    if value is None or isinstance(value, bool):
        return False
    try:
        money(str(value))
    except (ValueError, TypeError, ArithmeticError):
        return False
    return True


def _rows(path):
    from openpyxl import load_workbook

    report = inspect_workbook(path)
    receipts, products = [], []
    source_rows = []
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet_index, sheet in enumerate(book, 1):
            sheet.reset_dimensions()
            values = sheet.iter_rows(values_only=True)
            headers = _headers(next(values, ()))
            profile = ('bonsy' if {'Eintrags-ID', 'Anbieter', 'Datum & Uhrzeit', 'Summe'} <= set(headers)
                       else 'products' if {'Eintrags-ID', 'Produktname/Pfand', 'bezahlter Preis'} <= set(headers)
                       else None)
            if profile is None:
                continue
            nonempty_headers = [header for header in headers if header]
            if len(nonempty_headers) != len(set(nonempty_headers)):
                raise ValueError('invalid_bonsy_source_rows')
            for number, cells in enumerate(values, 2):
                if not any(value is not None for value in cells):
                    continue
                if len(receipts) + len(products) >= MAX_ROWS:
                    raise ValueError('too_many_bonsy_rows')
                row = dict(zip(headers, cells))
                row['_line'] = number
                row['_sheet_index'] = sheet_index
                source_rows.append((profile, row))
                (receipts if profile == 'bonsy' else products).append(row)
    finally:
        book.close()
    excluded_receipt_ids = {
        str(row.get('Eintrags-ID')).strip() for row in receipts
        if not _blank_source_value(row.get('Eintrags-ID'))
        and _text(row.get('Statistik-Info'), 'statistic_flag', 40, optional=True) == 'Exkludiert'
    }
    allowed_missing_excluded_fields = set()
    for profile, row in source_rows:
        is_excluded_receipt = (
            profile == 'bonsy'
            and _text(row.get('Statistik-Info'), 'statistic_flag', 40, optional=True) == 'Exkludiert')
        if is_excluded_receipt:
            missing_date = _blank_source_value(row.get('Datum & Uhrzeit'))
            missing_total = _blank_source_value(row.get('Summe'))
            if (missing_date or missing_total) and (
                    missing_date or _valid_preview_date(row.get('Datum & Uhrzeit'))) and (
                    missing_total or _valid_preview_amount(row.get('Summe'))):
                allowed_missing_excluded_fields.add((row['_sheet_index'], row['_line']))
            continue
        if profile == 'products':
            entry_id = str(row.get('Eintrags-ID')).strip() if not _blank_source_value(
                row.get('Eintrags-ID')) else ''
            product_excluded = (
                entry_id in excluded_receipt_ids
                or _text(row.get('Statistik-Info'), 'statistic_flag', 40, optional=True) == 'Exkludiert')
            if product_excluded and _blank_source_value(row.get('bezahlter Preis')):
                allowed_missing_excluded_fields.add((row['_sheet_index'], row['_line']))
    for sheet in report['sheets']:
        if sheet['profile'] not in {'bonsy', 'products'}:
            continue
        issues = sheet['issues']
        # A truncated preview cannot prove that every issue is the narrow exception below.
        if sheet['issue_count'] != len(issues):
            raise ValueError('invalid_bonsy_source_rows')
        for issue in issues:
            if (issue['code'] != 'invalid_required_value'
                    or (sheet['sheet_index'], issue['row']) not in allowed_missing_excluded_fields):
                raise ValueError('invalid_bonsy_source_rows')
    if not receipts:
        raise ValueError('no_bonsy_receipts')
    return receipts, products


def _normalize(path):
    receipt_rows, product_rows = _rows(path)
    receipts, seen, excluded_receipts = [], set(), set()
    for row in receipt_rows:
        entry = _text(row.get('Eintrags-ID'), 'entry_id', 240)
        if entry in seen:
            raise ValueError('duplicate_bonsy_entry')
        seen.add(entry)
        if _text(row.get('Statistik-Info'), 'statistic_flag', 40, optional=True) == 'Exkludiert':
            excluded_receipts.add(entry)
            continue
        source_currency = _text(row.get('Währung') or 'EUR', 'currency', 3)
        if source_currency not in {'EUR', '€'}:
            raise ValueError('unsupported_bonsy_currency')
        currency = 'EUR'
        total_value = money(_decimal(row.get('Summe'), 'total'))
        if total_value == 0:
            raise ValueError('invalid_bonsy_total')
        total = format(total_value, '.2f')
        receipts.append({
            'entry_id': entry, 'occurred_at': _day(row.get('Datum & Uhrzeit')),
            'vendor': _text(row.get('Anbieter'), 'vendor', 240),
            'branch': _text(row.get('Branche'), 'branch', 240, optional=True),
            'total': total, 'currency': currency, 'is_refund': total_value < 0,
        })
    products = defaultdict(list)
    excluded_products = []
    imported_receipt_ids = {receipt['entry_id'] for receipt in receipts}
    for row in product_rows:
        entry = _text(row.get('Eintrags-ID'), 'entry_id', 240)
        if entry not in seen:
            raise ValueError('orphan_bonsy_product')
        line_number = len(products[entry]) + 1
        if (entry in excluded_receipts
                or _text(row.get('Statistik-Info'), 'statistic_flag', 40, optional=True) == 'Exkludiert'):
            excluded_products.append(f'{entry}:{row["_line"]}')
            continue
        if entry not in imported_receipt_ids:
            continue
        products[entry].append({
            'line_number': line_number,
            'product_name': _text(row.get('Produktname/Pfand'), 'product_name', 500),
            'source_category': _text(row.get('Kategorie'), 'product_category', 240, optional=True),
            'paid_price': _decimal(row.get('bezahlter Preis'), 'paid_price', optional=True),
            'quantity': _decimal(row.get('Menge'), 'quantity', optional=True),
            'unit': _text(row.get('Einheit'), 'unit', 80, optional=True),
            'unit_price': _decimal(row.get('Preis (pro Stück/kg)'), 'unit_price', optional=True),
            'discount': _decimal(row.get('Rabatt'), 'discount', optional=True),
            'discount_name': _text(row.get('Name des Rabatts'), 'discount_name', 240, optional=True),
        })
    return receipts, products, {
        'excluded_receipt_ids': sorted(excluded_receipts),
        'excluded_product_ids': excluded_products,
        'source_inspected': True,
    }


def _title(receipt, products):
    names = [item['product_name'] for item in products[:3]]
    detail = ', '.join(names)
    if len(products) > 3:
        detail += f' und {len(products) - 3} weitere Positionen'
    document_type = 'Gutschrift' if receipt['is_refund'] else 'Kassenbon'
    return f"{document_type} · {detail}" if detail else f"{document_type} · {receipt['vendor']}"


def _source_text(receipt, products):
    lines = [f"Anbieter: {receipt['vendor']}", f"Datum: {receipt['occurred_at']}",
             f"Gesamt: {receipt['total']} EUR"]
    if receipt['branch']:
        lines.append(f"Branche: {receipt['branch']}")
    if products:
        lines.append('Positionen:')
        for item in products:
            price = f" · {item['paid_price']} EUR" if item['paid_price'] is not None else ''
            category = f" · Quelle: {item['source_category']}" if item['source_category'] else ''
            lines.append(f"- {item['product_name']}{price}{category}")
    return '\n'.join(lines)


def _reconcile_excluded_receipts(store, entry_ids):
    """Withdraw local Bonsy annotations only for exclusions parsed from this source."""
    result = {'reconciled_receipt_ids': [], 'unchanged_receipt_ids': [],
              'missing_receipt_ids': []}
    for entry_id in sorted(set(entry_ids)):
        document = store.db.execute(
            'SELECT d.* FROM bonsy_receipts r '
            'JOIN classification_documents d ON d.id=r.document_id '
            'WHERE r.entry_id=?', (entry_id,)).fetchone()
        if document is None:
            result['missing_receipt_ids'].append(entry_id)
            continue
        if document['source_reference'] != 'bonsy:' + entry_id:
            raise ValueError('bonsy_receipt_document_mismatch')
        links = store.db.execute(
            'SELECT account_id,external_id,document_id,allocated_amount,allocation_type '
            'FROM classification_document_links WHERE document_id=? ORDER BY account_id,external_id',
            (document['id'],)).fetchall()
        allocations = store.db.execute(
            'SELECT entry_id,account_id,external_id,allocated_amount,confirmed_at '
            'FROM bonsy_cash_allocations WHERE entry_id=? ORDER BY account_id,external_id',
            (entry_id,)).fetchall()
        warnings = json.loads(document['warnings'])
        if not isinstance(warnings, list) or any(not isinstance(item, str) for item in warnings):
            raise ValueError('invalid_bonsy_document_warnings')
        if (document['status'] == 'unreviewed' and _EXCLUDED_BONSY_WARNING in warnings
                and not links and not allocations):
            result['unchanged_receipt_ids'].append(entry_id)
            continue

        previous = dict(document)
        current = previous | {
            'status': 'unreviewed',
            'warnings': json.dumps(sorted(set(warnings) | {_EXCLUDED_BONSY_WARNING})),
            'revision': previous['revision'] + 1,
        }
        _audit(store, 'bonsy_source_excluded', current, document_id=document['id'],
               previous=previous)
        store.db.execute(
            'UPDATE classification_documents SET status=?,warnings=?,revision=? WHERE id=?',
            (current['status'], current['warnings'], current['revision'], document['id']))
        for link in links:
            link = dict(link)
            _audit(store, 'bonsy_excluded_link_removed', {'removed': True},
                   account_id=link['account_id'], external_id=link['external_id'],
                   document_id=document['id'], previous=link)
            store.db.execute(
                'DELETE FROM classification_document_links '
                'WHERE account_id=? AND external_id=? AND document_id=?',
                (link['account_id'], link['external_id'], document['id']))
        for allocation in allocations:
            allocation = dict(allocation)
            _audit(store, 'bonsy_excluded_cash_allocation_removed', {'removed': True},
                   account_id=allocation['account_id'], external_id=allocation['external_id'],
                   document_id=document['id'], previous=allocation)
            store.db.execute(
                'DELETE FROM bonsy_cash_allocations WHERE entry_id=? AND account_id=? AND external_id=?',
                (entry_id, allocation['account_id'], allocation['external_id']))
        result['reconciled_receipt_ids'].append(entry_id)
    return result


def import_workbook(store, database, path, *, confirm_exclusions=False):
    """Import structured receipts and products, then link only unique direct matches."""
    if type(confirm_exclusions) is not bool:
        raise ValueError('invalid_confirm_exclusions')
    path = outside_repository(path)
    if not path.is_file() or path.stat().st_size > MAX_BYTES:
        raise ValueError('invalid_bonsy_workbook')
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    receipts, products, exclusion_report = _normalize(path)
    previous = store.db.execute('SELECT * FROM bonsy_imports WHERE source_sha256=?',
                                (source_hash,)).fetchone()
    if previous is not None:
        reconciliation = {'reconciled_receipt_ids': [], 'unchanged_receipt_ids': [],
                          'missing_receipt_ids': []}
        if confirm_exclusions:
            store.db.execute('BEGIN IMMEDIATE')
            try:
                reconciliation = _reconcile_excluded_receipts(
                    store, exclusion_report['excluded_receipt_ids'])
                store.db.commit()
            except Exception:
                store.db.rollback()
                raise
        excluded_ids = set(exclusion_report['excluded_receipt_ids'])
        source_document_ids = [row['document_id'] for row in store.db.execute(
            'SELECT entry_id,document_id FROM bonsy_receipts WHERE source_sha256=? ORDER BY document_id',
            (source_hash,)) if row['entry_id'] not in excluded_ids]
        repaired = auto_link_timestamped_receipts(store, source_document_ids)
        return {'status': 'duplicate', 'source_sha256': source_hash,
                'receipts': previous['receipt_count'], 'products': previous['product_count'],
                'new_receipts': 0, 'confirmed_receipts': 0,
                'auto_links': repaired, 'exclusion_report': exclusion_report,
                'reconciliation': reconciliation}
    reconciliation = {'reconciled_receipt_ids': [], 'unchanged_receipt_ids': [],
                      'missing_receipt_ids': []}
    new_documents = []
    store.db.execute('BEGIN IMMEDIATE')
    try:
        store.db.execute('INSERT INTO bonsy_imports VALUES (?,?,?,?)', (
            source_hash, datetime.now(UTC).isoformat(), len(receipts),
            sum(len(items) for items in products.values())))
        for receipt in receipts:
            existing = store.db.execute('SELECT * FROM bonsy_receipts WHERE entry_id=?',
                                        (receipt['entry_id'],)).fetchone()
            if existing is not None:
                expected = (receipt['occurred_at'], receipt['vendor'], receipt['branch'],
                            receipt['total'], receipt['currency'])
                actual = tuple(existing[key] for key in
                               ('occurred_at', 'vendor', 'branch', 'total', 'currency'))
                if actual != expected:
                    raise ValueError('conflicting_bonsy_entry')
                continue
            item_products = products.get(receipt['entry_id'], [])
            document_id = store.db.execute(
                """INSERT INTO classification_documents
                   (kind,vendor,title,document_date,amount,currency,source_reference,warnings,status,revision)
                   VALUES ('invoice',?,?,?,?,?,?, '[]','confirmed',1)""",
                (receipt['vendor'], _title(receipt, item_products),
                 receipt['occurred_at'][:10], format(abs(money(receipt['total'])), '.2f'), 'EUR',
                 'bonsy:' + receipt['entry_id'])).lastrowid
            store.db.execute(
                'INSERT INTO bonsy_receipts VALUES (?,?,?,?,?,?,?,?)',
                (receipt['entry_id'], document_id, source_hash, receipt['occurred_at'],
                 receipt['vendor'], receipt['branch'], receipt['total'], receipt['currency']))
            store.db.executemany(
                'INSERT INTO bonsy_products VALUES (?,?,?,?,?,?,?,?,?,?)',
                [(receipt['entry_id'], item['line_number'], item['product_name'],
                  item['source_category'], item['paid_price'], item['quantity'], item['unit'],
                  item['unit_price'], item['discount'], item['discount_name'])
                 for item in item_products])
            new_documents.append((document_id, _source_text(receipt, item_products),
                                  receipt['is_refund']))
        if confirm_exclusions:
            reconciliation = _reconcile_excluded_receipts(
                store, exclusion_report['excluded_receipt_ids'])
        store.db.commit()
    except Exception:
        store.db.rollback()
        raise
    for document_id, text, _ in new_documents:
        cache_document_source(database, document_id, text)
    # A negative Bonsy total identifies a refund, but the generic matcher can
    # also consider same-value debits. Keep these reviewable instead of risking
    # an automatic link in the wrong direction.
    ids = [document_id for document_id, _, is_refund in new_documents if not is_refund]
    auto_links = auto_link_documents(store, {'confirmed': True, 'ids': ids})
    timestamp_links = auto_link_timestamped_receipts(store, ids)
    auto_links['linked'].extend(timestamp_links['linked'])
    auto_links['rejected'].extend(timestamp_links['rejected'])
    auto_links['checked'] += timestamp_links['checked']
    return {'status': 'ok', 'source_sha256': source_hash, 'receipts': len(receipts),
            'products': sum(len(items) for items in products.values()),
            'new_receipts': len(new_documents), 'confirmed_receipts': len(new_documents),
            'auto_links': auto_links,
            'unlinked_receipts': len(new_documents) - len(auto_links['linked']),
            'exclusion_report': exclusion_report, 'reconciliation': reconciliation}
