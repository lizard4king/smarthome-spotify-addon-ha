"""Read-only XLSX source inspection. Never writes ledger entries or evaluates formulas."""
from collections import Counter
from datetime import date, datetime
import hashlib
import json
import re
import warnings
import zipfile

from .core import money
from .storage_paths import (
    RepositoryPathError as RepositoryPathError,
    outside_repository as outside_repository,
)

MAX_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 100

PROFILES = {
    'finanzguru': {'Buchungstag', 'Referenzkonto', 'Name Referenzkonto', 'Betrag', 'Waehrung'},
    'bonsy': {'Eintrags-ID', 'Anbieter', 'Datum & Uhrzeit', 'Summe', 'Statistik-Info'},
    'products': {'Eintrags-ID', 'Produktname/Pfand', 'Kategorie', 'bezahlter Preis'},
}


def _day(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return datetime.fromisoformat(value).date().isoformat()
    raise ValueError('Invalid date')


def _amount(value):
    # XLSX numeric cells are decoded by openpyxl. Do not round sub-cent values.
    if value is None or isinstance(value, bool):
        raise ValueError('Missing amount')
    return money(str(value))


def _headers(row):
    return [str(value).strip() if value is not None else '' for value in row]


def inspect_workbook(path):
    """Return counts and diagnostics only; no account names or transaction descriptions."""
    from openpyxl import load_workbook
    from defusedxml.ElementTree import iterparse

    path = outside_repository(path)
    if path.stat().st_size > MAX_BYTES:
        raise ValueError('Quelldatei überschreitet 32 MiB.')
    if not zipfile.is_zipfile(path):
        raise ValueError('Kein XLSX-Container; alte XLS-Dateien werden nicht unterstützt.')
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10_000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
            raise ValueError('XLSX-Inhalt überschreitet die Größenbegrenzung.')
        if any('vbaproject' in e.filename.lower() for e in entries):
            raise ValueError('Makro-Inhalte werden nicht unterstützt.')
        total_rows = 0
        for entry in entries:
            if not re.fullmatch(r'xl/worksheets/sheet\d+\.xml', entry.filename):
                continue
            with archive.open(entry) as xml:
                for _, element in iterparse(xml, events=('end',)):
                    tag = element.tag.rsplit('}', 1)[-1]
                    if tag == 'row':
                        total_rows += 1
                        if total_rows > MAX_ROWS:
                            raise ValueError('Zu viele Zeilen im Workbook.')
                    if tag == 'c':
                        match = re.fullmatch(r'([A-Z]+)([1-9][0-9]*)', element.get('r', ''))
                        if not match:
                            raise ValueError('Ungültige Zelladresse.')
                        column = 0
                        for letter in match[1]:
                            column = column * 26 + ord(letter) - ord('A') + 1
                        if column > MAX_COLUMNS or int(match[2]) > MAX_ROWS:
                            raise ValueError('Zelladresse außerhalb der Vorschaugrenze.')
                    element.clear()
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    # A connector may give an XLSX file an XLS extension. Inspect the bytes.
    with path.open('rb') as stream, warnings.catch_warnings():
        warnings.filterwarnings('ignore', message="Workbook contains no default style.*")
        workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            return _inspect(workbook, digest)
        finally:
            workbook.close()


def _inspect(workbook, digest):
    result = {'source_sha256': digest, 'format': 'xlsx', 'ledger_written': False,
              'sheets': [], 'warnings': [], 'ready_for_ledger': False}
    receipt_ids, product_ids = set(), set()
    for sheet_index, sheet in enumerate(workbook, 1):
        if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ValueError('Tabellengröße überschreitet die Vorschaugrenze.')
        # Producers may omit or misstate dimensions; inspect actual cells.
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_col=MAX_COLUMNS + 1)
        first = next(rows, ())
        header = _headers([c.value for c in first])
        while header and not header[-1]:
            header.pop()
        if len(header) > MAX_COLUMNS:
            raise ValueError('Zu viele Spalten.')
        profile = next((name for name, required in PROFILES.items()
                        if required <= set(header)), 'planning_or_unknown')
        summary = {'sheet_index': sheet_index, 'profile': profile, 'rows': 0,
                   'valid_rows': 0, 'formula_cells': 0, 'issues': [],
                   'issue_count': 0, 'duplicate_ids': 0, 'conflicting_ids': 0}
        result['sheets'].append(summary)

        def issue(row_number, code):
            summary['issue_count'] += 1
            if len(summary['issues']) < 100:
                summary['issues'].append({'row': row_number, 'code': code})

        nonempty = [x for x in header if x]
        if len(nonempty) != len(set(nonempty)):
            issue(1, 'duplicate_headers')
            profile = summary['profile'] = 'planning_or_unknown'
        ids, dates, currencies, exclusions = {}, [], Counter(), Counter()
        split_rows = []
        for row_number, cells in enumerate(rows, 2):
            if row_number > MAX_ROWS:
                raise ValueError('Zu viele Zeilen.')
            if cells[-1].value is not None:
                raise ValueError('Zu viele Spalten.')
            if not any(c.value is not None for c in cells):
                continue
            summary['rows'] += 1
            formulas = sum(c.data_type == 'f' for c in cells)
            summary['formula_cells'] += formulas
            if profile == 'planning_or_unknown':
                continue
            if formulas:
                issue(row_number, 'formula_in_source_row')
                continue
            values = [c.value for c in cells[:len(header)]]
            row = dict(zip(header, values))
            if profile == 'finanzguru' and row.get('Split-Typ'):
                split_rows.append(row)
            issues_before = summary['issue_count']
            try:
                if profile == 'finanzguru':
                    day = _day(row['Buchungstag'])
                    _amount(row['Betrag'])
                    if not row['Referenzkonto']:
                        raise ValueError('Missing account')
                    currency = row['Waehrung']
                    external_id = row.get('Buchungs-ID')
                    key = (str(row['Referenzkonto']), str(external_id))
                    if row.get('Split-Typ'):
                        exclusions['split_rows_requiring_review'] += 1
                elif profile == 'bonsy':
                    day = _day(row['Datum & Uhrzeit'])
                    _amount(row['Summe'])
                    currency = row.get('Währung', row.get('Waehrung'))
                    external_id = row['Eintrags-ID']
                    key = str(external_id)
                    flag = row['Statistik-Info']
                    exclusions[flag if flag in {'Inkludiert', 'Exkludiert'} else 'unknown_statistic_flag'] += 1
                else:
                    _amount(row['bezahlter Preis'])
                    external_id = row['Eintrags-ID']
                    if not external_id:
                        raise ValueError('Missing receipt ID')
                    product_ids.add(str(external_id))
                    summary['valid_rows'] += 1
                    continue
                if currency not in {'EUR', '€'}:
                    issue(row_number, 'unsupported_currency')
                    continue
                if not external_id:
                    issue(row_number, 'missing_source_id')
                else:
                    if profile == 'bonsy':
                        receipt_ids.add(str(external_id))
                    fingerprint = hashlib.sha256(json.dumps(values, default=str,
                                                           ensure_ascii=False).encode()).hexdigest()
                    if key in ids:
                        if ids[key] == fingerprint:
                            summary['duplicate_ids'] += 1
                        else:
                            summary['conflicting_ids'] += 1
                            issue(row_number, 'conflicting_source_id')
                    else:
                        ids[key] = fingerprint
                dates.append(day)
                currencies['EUR'] += 1
                # This is value validation only, not permission to import.
                summary['valid_rows'] += int(summary['issue_count'] == issues_before)
            except (ValueError, TypeError, KeyError, ArithmeticError):
                issue(row_number, 'invalid_required_value')
        if dates:
            summary['date_from'], summary['date_to'] = min(dates), max(dates)
        summary['currencies'] = dict(currencies)
        summary['source_flags'] = dict(exclusions)
        if profile == 'finanzguru':
            from .split_preview import inspect_split_groups
            summary['splits'] = inspect_split_groups(split_rows)
    if product_ids:
        result['unmatched_product_receipt_ids'] = len(product_ids - receipt_ids)
    if not any(s['profile'] != 'planning_or_unknown' for s in result['sheets']):
        result['warnings'].append('no_supported_source_table')
    result['warnings'].append('preview_only_no_ownership_or_payment_matching')
    return result
