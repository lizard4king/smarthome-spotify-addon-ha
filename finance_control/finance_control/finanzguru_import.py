"""Private staged XLSX intake; explicit account and transfer decisions, atomic ledger import."""
import csv
import hashlib
import io
import json
import re
import secrets
from collections import Counter
from datetime import UTC, date, datetime

from .core import money
from .classification import apply_exact_rule_if_safe, exact_rule_match
from .import_preview import (
    MAX_BYTES,
    _amount,
    _day,
    _headers,
    inspect_workbook,
    outside_repository,
)
from .split_preview import inspect_split_groups

MAX_STAGE_ROWS = 20_000
HEADER = ['external_id', 'account_id', 'date', 'amount', 'currency', 'category', 'transfer_id']


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def stage_metadata_digest(staged):
    """Bind every staged import decision to the archived source review."""
    return digest({key: staged.get(key) for key in (
        'source_sha256', 'rows', 'blockers', 'split_summary', 'split_row_count',
        'excluded_split_children',
    )})


def private_directory(root):
    root = outside_repository(root)
    root.mkdir(parents=True, exist_ok=True)
    return root


def source_category(row):
    """Keep the export's category label without inventing a Finance-Control category."""
    category = row.get('Kategorie')
    if category is not None and str(category).strip():
        return str(category).strip()
    analysis = [str(row.get(column)).strip() for column in
                ('Analyse-Hauptkategorie', 'Analyse-Unterkategorie')
                if row.get(column) is not None and str(row.get(column)).strip()]
    return ' / '.join(analysis) if analysis else 'Unkategorisiert (Quelle)'


def source_text(row, column, maximum):
    """Preserve optional booking context without guessing from account labels."""
    value = row.get(column)
    if value is None:
        return ''
    value = str(value).strip()
    if len(value) > maximum:
        raise ValueError('source_text_too_long')
    return value


def stage_workbook(raw, root):
    """Archive exact source bytes and normalized rows outside Git. No ledger access."""
    from openpyxl import load_workbook

    if not isinstance(raw, bytes) or not 0 < len(raw) <= MAX_BYTES:
        raise ValueError('invalid_workbook_size')
    root = private_directory(root)
    batch = secrets.token_hex(16)
    folder = root / batch
    folder.mkdir()
    source = folder / 'source.xlsx'
    source.write_bytes(raw)
    report = inspect_workbook(source)
    supported = [s for s in report['sheets'] if s['profile'] == 'finanzguru']
    blockers = []
    if not supported:
        blockers.append('no_finanzguru_sheet')
    if sum(s['rows'] for s in supported) > MAX_STAGE_ROWS:
        blockers.append('too_many_rows')
    for sheet in supported:
        if sheet['issue_count']:
            blockers.append('invalid_source_rows')
    rows = []
    split_rows = []
    excluded_split_children = 0
    split_date_mismatches = 0
    source_id_kinds = {}
    # Inspect all supported rows even when the preflight has findings: split-child
    # identities must not escape validation merely because their parent is blocked.
    if supported:
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        try:
            indexes = {s['sheet_index'] for s in supported}
            for index, sheet in enumerate(book, 1):
                if index not in indexes:
                    continue
                sheet.reset_dimensions()
                values = sheet.iter_rows(values_only=True)
                headers = _headers(next(values, ()))
                for number, cells in enumerate(values, 2):
                    if not any(v is not None for v in cells):
                        continue
                    row = dict(zip(headers, cells))
                    account, identifier = row['Referenzkonto'], row.get('Buchungs-ID')
                    if (not isinstance(account, str) or not account.strip()
                            or not isinstance(identifier, str) or not identifier.strip()):
                        blockers.append('source_ids_must_be_text')
                        continue
                    kind = row.get('Split-Typ')
                    source_id_kinds.setdefault((account, identifier), []).append(kind)
                    if kind:
                        split_rows.append(row)
                    if kind in {'Teilbuchung', 'Restbetrag'}:
                        parent_id = row.get('Referenz-Original-ID')
                        if not isinstance(parent_id, str) or not parent_id.strip():
                            blockers.append('source_ids_must_be_text')
                            continue
                        excluded_split_children += 1
                        continue
                    if kind and kind != 'Original':
                        # The group will be blocked below. Do not mistake an unknown split row
                        # for an ordinary ledger transaction while it is under review.
                        continue
                    try:
                        rows.append({'reference': f'{index}:{number}', 'source_account': account,
                                     'source_name': str(row.get('Name Referenzkonto') or ''),
                                     'source_id': identifier,
                                     'external_id': 'fg:' + digest(account) + ':' + digest(identifier),
                                     'date': _day(row['Buchungstag']),
                                     'amount': format(_amount(row['Betrag']), '.2f'), 'currency': 'EUR',
                                     'category': source_category(row),
                                     'counterparty': source_text(
                                         row, 'Beguenstigter/Auftraggeber', 240),
                                     'description': source_text(row, 'Verwendungszweck', 1000)})
                    except (ValueError, TypeError, KeyError, ArithmeticError):
                        blockers.append('invalid_source_rows')
        finally:
            book.close()
    split_summary = inspect_split_groups(split_rows)
    split_parents = {}
    for row in split_rows:
        account, identifier, kind = row.get('Referenzkonto'), row.get('Buchungs-ID'), row.get('Split-Typ')
        if kind == 'Original':
            split_parents[(account, identifier)] = row
    for child in split_rows:
        if child.get('Split-Typ') in {'Teilbuchung', 'Restbetrag'}:
            parent = split_parents.get((child.get('Referenzkonto'), child.get('Referenz-Original-ID')))
            if parent is not None and child.get('Buchungstag') != parent.get('Buchungstag'):
                split_date_mismatches += 1
    split_id_collisions = sum(
        len(kinds) - 1 for kinds in source_id_kinds.values()
        if len(kinds) > 1 and any(kind in {'Teilbuchung', 'Restbetrag'} for kind in kinds)
    )
    split_summary = split_summary | {'date_mismatches': split_date_mismatches,
                                     'source_id_collisions': split_id_collisions}
    split_blockers = {
        'invalid_split_groups': split_summary['invalid_groups'],
        'unknown_split_rows': split_summary['unknown_split_rows'],
        'orphan_or_ambiguous_split_groups': split_summary['missing_or_ambiguous_original'],
        'unbalanced_split_groups': split_summary['unbalanced'],
        'split_date_mismatch': split_date_mismatches,
        'split_source_id_collision': split_id_collisions,
    }
    blockers.extend(name for name, count in split_blockers.items() if count)
    # Do not leave a partially usable stage behind after a source or split validation failure.
    if blockers:
        rows = []
        excluded_split_children = 0
    staged = {'version': 2, 'source_sha256': hashlib.sha256(raw).hexdigest(), 'rows_sha256': digest(rows),
              'rows': rows, 'blockers': sorted(set(blockers)), 'format_report': report,
              'split_summary': split_summary, 'split_row_count': len(split_rows),
              'excluded_split_children': excluded_split_children}
    staged['metadata_sha256'] = stage_metadata_digest(staged)
    (folder / 'stage.json').write_text(json.dumps(staged, ensure_ascii=False), encoding='utf-8')
    return {'batch': batch, 'blockers': staged['blockers'], 'source_sha256': staged['source_sha256'],
            'rows': rows, 'source_rows': sum(s['rows'] for s in supported),
            'issues': [{'sheet': s['sheet_index'], **issue} for s in supported for issue in s['issues']],
            'ignored_sheets': len(report['sheets']) - len(supported),
            'split_summary': split_summary, 'excluded_split_children': excluded_split_children,
            'ledger_written': False}


def load_batch(root, batch):
    if not isinstance(batch, str) or not re.fullmatch('[0-9a-f]{32}', batch):
        raise ValueError('invalid_batch')
    root = outside_repository(root)
    folder = outside_repository(root / batch)
    if not folder.is_relative_to(root):
        raise ValueError('redirected_batch')
    staged = json.loads((folder / 'stage.json').read_text(encoding='utf-8'))
    if hashlib.sha256((folder / 'source.xlsx').read_bytes()).hexdigest() != staged['source_sha256']:
        raise ValueError('source_changed')
    if digest(staged['rows']) != staged['rows_sha256']:
        raise ValueError('staged_rows_changed')
    if staged.get('version', 1) >= 2 and staged.get('metadata_sha256') != stage_metadata_digest(staged):
        raise ValueError('staged_metadata_changed')
    return folder, staged


def prepare(store, staged, choices):
    """Evaluate decisions against the current ledger; never guesses ownership/transfers."""
    accounts = {a['id']: a for a in store.accounts()}
    mapping = choices.get('mapping', {})
    pairs = choices.get('transfer_pairs', [])
    if not isinstance(mapping, dict) or not isinstance(pairs, list):
        raise TypeError('invalid_choices')
    if choices.get('transfers_reviewed') is not True:
        raise ValueError('transfer_review_required')
    all_rows = staged['rows']
    date_to = choices.get('date_to')
    if date_to is not None:
        if not isinstance(date_to, str) or not re.fullmatch(r'[0-9]{4}-[0-9]{2}-[0-9]{2}', date_to):
            raise ValueError('invalid_date_to')
        try:
            date_to = date.fromisoformat(date_to).isoformat()
        except ValueError as error:
            raise ValueError('invalid_date_to') from error
    rows = [row for row in all_rows if date_to is None or row['date'] <= date_to]
    excluded_after_date = [row for row in all_rows if date_to is not None and row['date'] > date_to]
    blockers = list(staged['blockers'])
    if staged.get('split_row_count', 0) and choices.get('split_policy') != 'originals_only':
        blockers.append('split_policy_required')
    source_accounts = {r['source_account'] for r in all_rows}
    managed_sources = {row[0] for row in store.db.execute(
        'SELECT source_account FROM finanzguru_snapshot_heads')}
    if source_accounts & managed_sources:
        blockers.append('snapshot_managed_source_requires_full_export')
    if set(mapping) != source_accounts or any(target not in accounts for target in mapping.values()):
        raise ValueError('complete_account_mapping_required')
    if len(set(mapping.values())) != len(mapping):
        raise ValueError('source_accounts_must_map_to_distinct_targets')
    references = {r['reference']: r for r in rows}
    transfer_ids = {}
    for pair in pairs:
        if (not isinstance(pair, list) or len(pair) != 2 or pair[0] == pair[1]
                or any(ref not in references or ref in transfer_ids for ref in pair)):
            raise ValueError('invalid_transfer_pair')
        a, b = (references[ref] for ref in pair)
        if (mapping[a['source_account']] == mapping[b['source_account']]
                or money(a['amount']) == 0 or money(a['amount']) + money(b['amount']) != 0):
            raise ValueError('invalid_transfer_amounts')
        if any(r['date'] <= accounts[mapping[r['source_account']]]['opening_date'] for r in (a, b)):
            raise ValueError('transfer_crosses_opening_date')
        identity = 'fg-transfer:' + digest(sorted([a['external_id'], b['external_id']]))
        for ref in pair:
            transfer_ids[ref] = identity
    ledger = [tuple(row) for row in store.db.execute(
        'SELECT account_id,external_id,date,amount,currency,category,transfer_id FROM transactions ORDER BY account_id,external_id')]
    existing = {(r[0], r[1]): r for r in ledger}
    existing_elsewhere = {r[1]: r[0] for r in ledger if r[1].startswith('fg:')}
    account_bindings = {}
    for existing_row in ledger:
        if existing_row[1].startswith('fg:'):
            prefix = existing_row[1].rsplit(':', 1)[0]
            account_bindings.setdefault(prefix, set()).add(existing_row[0])
    by_date_amount = {(r[0], r[2], r[3]) for r in ledger}
    seen, detail, proposed, import_rows = {}, [], [], []
    counts = Counter({'after_date_to': len(excluded_after_date)}) if excluded_after_date else Counter()
    for row in excluded_after_date:
        detail.append({'reference': row['reference'], 'account_id': mapping[row['source_account']],
                       'date': row['date'], 'amount': row['amount'], 'category': row['category'],
                       'status': 'after_date_to', 'transfer': False})
    for row in rows:
        target = mapping[row['source_account']]
        record = (target, row['external_id'], row['date'], row['amount'], 'EUR', row['category'],
                  transfer_ids.get(row['reference'], ''))
        key = record[:2]
        binding = account_bindings.get(row['external_id'].rsplit(':', 1)[0], set())
        if binding and binding != {target}:
            status = 'account_mapping_conflict'
        elif row['date'] <= accounts[target]['opening_date']:
            status = 'before_or_on_opening'
        elif row['external_id'] in existing_elsewhere and existing_elsewhere[row['external_id']] != target:
            status = 'account_mapping_conflict'
        elif key in seen:
            status = 'duplicate_in_file' if seen[key] == record else 'source_id_conflict'
        elif key in existing:
            status = 'already_imported' if existing[key] == record else 'ledger_id_conflict'
        elif (target, row['date'], row['amount']) in by_date_amount:
            status = 'possible_other_import_overlap'
        else:
            status = 'new'
            proposed.append(record)
            import_rows.append((row, target, record))
        seen[key] = record
        counts[status] += 1
        if status in {'source_id_conflict', 'ledger_id_conflict', 'account_mapping_conflict', 'possible_other_import_overlap'}:
            blockers.append(status)
        direction = 'income' if money(row['amount']) > 0 else 'expense'
        rule = None if status != 'new' or record[-1] else exact_rule_match(
            store, row.get('counterparty', ''), row.get('description', ''), direction)
        detail.append({'reference': row['reference'], 'account_id': target, 'date': row['date'],
                       'amount': row['amount'], 'category': row['category'], 'status': status,
                       'transfer': bool(record[-1]),
                       'own_category': None if rule is None else rule['category'],
                       'own_category_label': None if rule is None else rule['label']})
    # Include existing pair legs as well, so the original importer revalidates the complete group.
    records = sorted({tuple(r) for r in seen.values() if r[2] > accounts[r[0]]['opening_date']})
    text = io.StringIO(newline='')
    writer = csv.writer(text, lineterminator='\n')
    writer.writerow(HEADER)
    for record in records:
        writer.writerow([record[1], record[0], *record[2:]])
    summary = {'counts': dict(counts), 'blockers': sorted(set(blockers)), 'rows': detail,
               'date_to': date_to,
               'auto_classification_count': sum(
                   row['status'] == 'new' and row['own_category'] is not None for row in detail),
               'new_net_amount': format(sum((money(r[3]) for r in proposed), money('0')), '.2f'),
               'ready': not blockers, 'ledger_written': False}
    token = digest({'source': staged['source_sha256'], 'choices': choices,
                    'rows': rows, 'blockers': staged['blockers'],
                    'split_summary': staged.get('split_summary', {}),
                    'split_row_count': staged.get('split_row_count', 0),
                    'excluded_split_children': staged.get('excluded_split_children', 0),
                    'accounts': list(accounts.values()), 'ledger': ledger})
    return summary | {'review_token': token}, text.getvalue(), import_rows


def preview_batch(store, root, batch, choices):
    if not isinstance(choices, dict) or choices.get('mode', 'append') not in {'append', 'snapshot'}:
        raise ValueError('invalid_import_mode')
    if choices.get('mode') == 'snapshot':
        from .finanzguru_snapshot import preview_snapshot
        return preview_snapshot(store, root, batch, choices)
    _, staged = load_batch(root, batch)
    return prepare(store, staged, choices)[0]


def commit_batch(store, root, batch, choices, review_token):
    if not isinstance(choices, dict) or choices.get('mode', 'append') not in {'append', 'snapshot'}:
        raise ValueError('invalid_import_mode')
    if choices.get('mode') == 'snapshot':
        from .finanzguru_snapshot import commit_snapshot
        return commit_snapshot(store, root, batch, choices, review_token)
    folder, staged = load_batch(root, batch)
    # Serialize review revalidation with all writes, including other local processes.
    store.db.execute('BEGIN IMMEDIATE')
    try:
        report, source, import_rows = prepare(store, staged, choices)
        if not report['ready'] or report['review_token'] != review_token:
            raise ValueError('review_stale_or_blocked')
        receipt = {'source_sha256': staged['source_sha256'], 'choices': choices,
                   'recorded_at': datetime.now(UTC).isoformat(),
                   'review_token': review_token, 'csv_sha256': hashlib.sha256(source.encode()).hexdigest()}
        # Durable source/decision before ledger commit; the imports table archives this exact CSV.
        with (folder / ('decision-' + secrets.token_hex(8) + '.json')).open('x', encoding='utf-8') as stream:
            json.dump(receipt, stream, ensure_ascii=False, indent=2)
        inserted = store.import_csv(source, manage_transaction=False)
        auto_classified = 0
        for row, account_id, record in import_rows:
            counterparty, description = row.get('counterparty', ''), row.get('description', '')
            store.db.execute(
                'INSERT INTO transaction_context VALUES (?,?,?,?)',
                (account_id, record[1], counterparty, description))
            auto_classified += apply_exact_rule_if_safe(
                store, account_id, record[1]) is not None
        store.db.commit()
    except Exception:
        store.db.rollback()
        raise
    # Optional enrichment happens after the atomic ledger commit. A failure
    # must never make a successful import look safe to repeat.
    from .transfer_corrections import auto_confirm_wallet_topups
    try:
        transfer_review = auto_confirm_wallet_topups(store)
    except Exception:
        transfer_review = {'status': 'failed', 'confirmed': 0, 'pair_ids': [],
                           'message': 'Automatische Aufladungen blieben zur Prüfung offen.'}
    from .local_model import review_transactions
    keys = [{'account_id': account_id, 'external_id': record[1]}
            for _, account_id, record in import_rows]
    try:
        model_review = review_transactions(store, keys)
    except Exception:
        # The ledger commit is already durable.  Never invite an unsafe retry
        # merely because optional local enrichment failed unexpectedly.
        model_review = {'status': 'failed', 'counts': {}, 'reviewed': 0,
                        'message': 'Lokale Modellprüfung fehlgeschlagen; Import blieb erhalten.'}
    return {'inserted': inserted, 'auto_classified': auto_classified,
            'auto_transfers': transfer_review,
            'model_review': model_review,
            'source_sha256': staged['source_sha256'], 'ledger_written': True}
